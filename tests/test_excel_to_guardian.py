import json
import os
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from src.cli import _load_plan, _write_plan
from src.excel_to_ir import ExcelValidationError, parse_guardian_excel
from src.ir_to_guardian import generate_script


USER_HEADERS = ["user_name", "email", "initial_password", "groups", "direct_roles"]
GROUP_HEADERS = ["group_name", "roles"]
PERMISSION_HEADERS = [
    "principal_type", "principal_name", "database", "table", "column", "path", "actions",
]


def make_workbook(path: Path, *, invalid: bool = False) -> None:
    workbook = Workbook()
    users = workbook.active
    users.title = "UserGroups"
    users.append(["填写用户、初始密码，以及用户所属组和直接角色。"])
    users.append(USER_HEADERS)
    users.append(["alice", "alice@example.com", "001234", "finance,audit", "direct_role"])

    groups = workbook.create_sheet("GroupRoles")
    groups.append(["填写组和角色关系；多个角色使用英文逗号分隔。"])
    groups.append(GROUP_HEADERS)
    groups.append(["finance", "finance_role,public_role"])

    permissions = workbook.create_sheet("Permissions")
    permissions.append(["根据 database/table/column/path 自动识别资源类型。"])
    permissions.append(PERMISSION_HEADERS)
    permissions.append(["ROLE", "finance_role", "GLOBAL", "", "", "", "ALL"])
    permissions.append(["USER", "alice", "sales", "orders", "", "", "SELECT,UPDATE"])
    permissions.append(["GROUP", "audit", "sales", "orders", "amount", "", "SELECT"])
    permissions.append(["USER", "alice", "", "", "", "/user/alice", "ALL"])
    if invalid:
        permissions.append(["USER", "undeclared", "sales", "", "", "", "OWNER"])
        permissions.append(["ROLE", "local_role", "", "", "", "file:///tmp", "READ"])

    workbook.save(path)


class ExcelToGuardianTests(unittest.TestCase):
    def test_excel_generates_entities_relationships_profiles_and_permissions(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            workbook_path = root / "guardian.xlsx"
            report_path = root / "validation.json"
            script_path = root / "guardian.sh"
            ir_path = root / "guardian.json"
            make_workbook(workbook_path)

            plan = parse_guardian_excel(
                str(workbook_path),
                validation_report=str(report_path),
            )
            _write_plan(plan, str(ir_path))
            restored = _load_plan(str(ir_path))
            generate_script(
                restored,
                str(script_path),
                base_url="https://guardian.example",
                access_token="unit-token",
                component_overrides={"hive": "unit-hive", "hdfs": "unit-hdfs"},
            )

            report = json.loads(report_path.read_text(encoding="utf-8"))
            script = script_path.read_text(encoding="utf-8")

        self.assertEqual(report["status"], "valid")
        self.assertEqual(plan.users, {"alice"})
        self.assertEqual(plan.groups, {"finance", "audit"})
        self.assertEqual(plan.roles, {"direct_role", "finance_role", "public_role"})
        self.assertEqual(plan.user_profiles["alice"]["initial_password"], "001234")
        self.assertEqual(plan.group_user_assignments["finance"], {"alice"})
        self.assertEqual(plan.role_user_assignments["direct_role"], {"alice"})
        self.assertEqual(plan.role_group_assignments["finance_role"], {"finance"})
        self.assertEqual(sum(len(item.permissions) for item in plan.policies), 15)
        self.assertIn('"userEmail": "alice@example.com"', script)
        self.assertIn('"userPassword": "001234"', script)
        self.assertIn('"component": "unit-hive"', script)
        self.assertIn('"component": "unit-hdfs"', script)
        self.assertIn('"principalType": "GROUP"', script)
        self.assertIn('["PATH", "/", "user", "alice"]', script)

    def test_excel_collects_all_errors_and_writes_report(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            workbook_path = root / "invalid.xlsx"
            report_path = root / "validation.json"
            make_workbook(workbook_path, invalid=True)

            with self.assertRaises(ExcelValidationError) as context:
                parse_guardian_excel(
                    str(workbook_path),
                    validation_report=str(report_path),
                )
            report = json.loads(report_path.read_text(encoding="utf-8"))

        self.assertEqual(report["status"], "errors")
        self.assertGreaterEqual(report["error_count"], 3)
        messages = [item["message"] for item in report["errors"]]
        self.assertTrue(any("UserGroups" in message for message in messages))
        self.assertTrue(any("OWNER" in message for message in messages))
        self.assertTrue(any("本地文件 URI" in message for message in messages))
        self.assertEqual(context.exception.report_path, os.path.abspath(report_path))


if __name__ == "__main__":
    unittest.main()
