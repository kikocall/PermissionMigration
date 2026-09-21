"""解析人工维护的 Guardian 批量权限 Excel 模板。"""

from __future__ import annotations

import json
import os
import re
from typing import Any, Optional

try:
    from .models import (
        MigrationPlan,
        PermissionEntry,
        Policy,
        Principal,
        PrincipalType,
        ResourcePath,
        ServiceType,
    )
except ImportError:
    from models import (  # type: ignore  # noqa
        MigrationPlan,
        PermissionEntry,
        Policy,
        Principal,
        PrincipalType,
        ResourcePath,
        ServiceType,
    )


USER_HEADERS = ["user_name", "email", "initial_password", "groups", "direct_roles"]
GROUP_HEADERS = ["group_name", "roles"]
PERMISSION_HEADERS = [
    "principal_type", "principal_name", "database", "table", "column", "path", "actions",
]

DATABASE_ACTIONS = ["CREATE", "SELECT", "INSERT", "UPDATE", "DELETE", "ADMIN", "ACCESS"]
TABLE_COLUMN_ACTIONS = ["SELECT", "INSERT", "UPDATE", "DELETE", "ADMIN"]
HDFS_ACTIONS = ["READ", "WRITE", "EXECUTE", "ADMIN", "ACCESS"]
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


class ExcelValidationError(ValueError):
    """Excel 全表校验失败。"""

    def __init__(self, errors: list[dict[str, Any]], report_path: Optional[str] = None):
        self.errors = errors
        self.report_path = report_path
        summary = "; ".join(
            f"{item['sheet']}!{item.get('row', '-')}/{item.get('field', '-')}: {item['message']}"
            for item in errors[:5]
        )
        if len(errors) > 5:
            summary += f"; 另有 {len(errors) - 5} 个错误"
        super().__init__(summary)


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _add_issue(
    issues: list[dict[str, Any]],
    sheet: str,
    row: int,
    field: str,
    message: str,
) -> None:
    issues.append({"sheet": sheet, "row": row, "field": field, "message": message})


def _parse_multi(
    value: str,
    *,
    sheet: str,
    row: int,
    field: str,
    errors: list[dict[str, Any]],
) -> list[str]:
    if not value:
        return []
    if "，" in value:
        _add_issue(errors, sheet, row, field, "只能使用英文逗号分隔多个值")
        return []
    items = [item.strip() for item in value.split(",")]
    if any(not item for item in items):
        _add_issue(errors, sheet, row, field, "列表中存在空项")
        return []
    return list(dict.fromkeys(items))


def _valid_name(
    value: str,
    *,
    sheet: str,
    row: int,
    field: str,
    errors: list[dict[str, Any]],
) -> bool:
    if not value:
        _add_issue(errors, sheet, row, field, "不能为空")
        return False
    if "," in value or "，" in value:
        _add_issue(errors, sheet, row, field, "名称中不能包含逗号")
        return False
    return True


def _headers(ws, expected: list[str], errors: list[dict[str, Any]]) -> bool:
    actual = [_cell_text(ws.cell(2, col).value) for col in range(1, len(expected) + 1)]
    if actual != expected:
        _add_issue(
            errors,
            ws.title,
            2,
            "headers",
            f"字段必须为 {expected}，实际为 {actual}",
        )
        return False
    return True


def _nonempty_rows(ws, width: int):
    for row_number in range(3, ws.max_row + 1):
        cells = [ws.cell(row_number, col) for col in range(1, width + 1)]
        if any(cell.value is not None and _cell_text(cell.value) for cell in cells):
            yield row_number, cells


def _write_report(path: str, report: dict[str, Any]) -> str:
    absolute = os.path.abspath(path)
    parent = os.path.dirname(absolute)
    if parent:
        os.makedirs(parent, exist_ok=True)
    with open(absolute, "w", encoding="utf-8") as stream:
        json.dump(report, stream, ensure_ascii=False, indent=2)
    return absolute


def parse_guardian_excel(
    filepath: str,
    *,
    validation_report: Optional[str] = None,
) -> MigrationPlan:
    """解析 Guardian Excel；任何校验错误都会阻止返回迁移计划。"""
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "读取 Excel 需要 openpyxl，请先执行: python -m pip install -r requirements.txt"
        ) from error

    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    plan = MigrationPlan()
    plan.source_metadata = {"source": "guardian-excel", "file": os.path.abspath(filepath)}
    row_counts = {"UserGroups": 0, "GroupRoles": 0, "Permissions": 0}
    case_names: dict[str, dict[str, str]] = {"USER": {}, "GROUP": {}, "ROLE": {}}

    def register_name(kind: str, name: str, sheet: str, row: int, field: str) -> None:
        if not name:
            return
        key = name.casefold()
        existing = case_names[kind].get(key)
        if existing is not None and existing != name:
            _add_issue(
                errors, sheet, row, field,
                f"{kind} 名称与已有名称仅大小写不同: {existing!r} / {name!r}",
            )
            return
        case_names[kind][key] = name

    try:
        workbook = load_workbook(filepath, data_only=False, read_only=False)
    except Exception as error:
        report = {
            "status": "errors",
            "source": os.path.abspath(filepath),
            "error_count": 1,
            "warning_count": 0,
            "errors": [{"sheet": "workbook", "row": 0, "field": "file", "message": str(error)}],
            "warnings": [],
        }
        report_path = _write_report(validation_report, report) if validation_report else None
        raise ExcelValidationError(report["errors"], report_path) from error

    required_sheets = ["UserGroups", "GroupRoles", "Permissions"]
    for sheet_name in required_sheets:
        if sheet_name not in workbook.sheetnames:
            _add_issue(errors, sheet_name, 0, "sheet", "缺少必需工作表")

    user_ws = workbook["UserGroups"] if "UserGroups" in workbook.sheetnames else None
    group_ws = workbook["GroupRoles"] if "GroupRoles" in workbook.sheetnames else None
    permission_ws = workbook["Permissions"] if "Permissions" in workbook.sheetnames else None

    if user_ws is not None and _headers(user_ws, USER_HEADERS, errors):
        for row_number, cells in _nonempty_rows(user_ws, len(USER_HEADERS)):
            row_counts["UserGroups"] += 1
            if any(cell.data_type == "f" for cell in cells):
                _add_issue(errors, "UserGroups", row_number, "row", "数据单元格不支持公式")
                continue
            user, email, password, group_text, role_text = [_cell_text(cell.value) for cell in cells]
            valid_user = _valid_name(
                user, sheet="UserGroups", row=row_number, field="user_name", errors=errors,
            )
            if not email:
                _add_issue(errors, "UserGroups", row_number, "email", "不能为空")
            elif not EMAIL_RE.fullmatch(email):
                _add_issue(errors, "UserGroups", row_number, "email", "邮箱格式不正确")
            if not password:
                _add_issue(errors, "UserGroups", row_number, "initial_password", "不能为空")
            elif cells[2].data_type == "n":
                warnings.append({
                    "sheet": "UserGroups",
                    "row": row_number,
                    "field": "initial_password",
                    "message": "密码是数值单元格；如需保留前导零，请在 Excel 中设置为文本",
                })

            groups = _parse_multi(
                group_text, sheet="UserGroups", row=row_number, field="groups", errors=errors,
            )
            direct_roles = _parse_multi(
                role_text, sheet="UserGroups", row=row_number, field="direct_roles", errors=errors,
            )
            for group in groups:
                _valid_name(group, sheet="UserGroups", row=row_number, field="groups", errors=errors)
                register_name("GROUP", group, "UserGroups", row_number, "groups")
            for role in direct_roles:
                _valid_name(role, sheet="UserGroups", row=row_number, field="direct_roles", errors=errors)
                register_name("ROLE", role, "UserGroups", row_number, "direct_roles")

            if valid_user:
                register_name("USER", user, "UserGroups", row_number, "user_name")
                profile = {"email": email, "initial_password": password}
                existing = plan.user_profiles.get(user)
                if existing is not None and existing != profile:
                    _add_issue(
                        errors, "UserGroups", row_number, "user_name",
                        "同一用户重复填写了不同的邮箱或初始密码",
                    )
                else:
                    plan.users.add(user)
                    plan.user_profiles[user] = profile
                    for group in groups:
                        plan.groups.add(group)
                        plan.group_user_assignments.setdefault(group, set()).add(user)
                    for role in direct_roles:
                        plan.roles.add(role)
                        plan.role_user_assignments.setdefault(role, set()).add(user)

    if group_ws is not None and _headers(group_ws, GROUP_HEADERS, errors):
        for row_number, cells in _nonempty_rows(group_ws, len(GROUP_HEADERS)):
            row_counts["GroupRoles"] += 1
            if any(cell.data_type == "f" for cell in cells):
                _add_issue(errors, "GroupRoles", row_number, "row", "数据单元格不支持公式")
                continue
            group, role_text = [_cell_text(cell.value) for cell in cells]
            valid_group = _valid_name(
                group, sheet="GroupRoles", row=row_number, field="group_name", errors=errors,
            )
            roles = _parse_multi(
                role_text, sheet="GroupRoles", row=row_number, field="roles", errors=errors,
            )
            for role in roles:
                _valid_name(role, sheet="GroupRoles", row=row_number, field="roles", errors=errors)
                register_name("ROLE", role, "GroupRoles", row_number, "roles")
            if valid_group:
                register_name("GROUP", group, "GroupRoles", row_number, "group_name")
                plan.groups.add(group)
                for role in roles:
                    plan.roles.add(role)
                    plan.role_group_assignments.setdefault(role, set()).add(group)

    seen_permissions: set[tuple[str, str, str, tuple[str, ...]]] = set()
    if permission_ws is not None and _headers(permission_ws, PERMISSION_HEADERS, errors):
        for row_number, cells in _nonempty_rows(permission_ws, len(PERMISSION_HEADERS)):
            row_counts["Permissions"] += 1
            if any(cell.data_type == "f" for cell in cells):
                _add_issue(errors, "Permissions", row_number, "row", "数据单元格不支持公式")
                continue
            values = [_cell_text(cell.value) for cell in cells]
            raw_type, principal_name, database, table, column, path, action_text = values
            principal_type = raw_type.upper()
            if principal_type not in {"USER", "GROUP", "ROLE"}:
                _add_issue(
                    errors, "Permissions", row_number, "principal_type",
                    "只允许 USER、GROUP、ROLE",
                )
                continue
            if not _valid_name(
                principal_name,
                sheet="Permissions",
                row=row_number,
                field="principal_name",
                errors=errors,
            ):
                continue
            register_name(principal_type, principal_name, "Permissions", row_number, "principal_name")
            if principal_type == "USER" and principal_name not in plan.users:
                _add_issue(
                    errors, "Permissions", row_number, "principal_name",
                    "USER 必须先在 UserGroups 中声明邮箱和初始密码",
                )
            elif principal_type == "GROUP":
                plan.groups.add(principal_name)
            elif principal_type == "ROLE":
                plan.roles.add(principal_name)

            resource: Optional[ResourcePath] = None
            allowed_actions: list[str] = []
            scope = ""
            if path:
                if database or table or column:
                    _add_issue(
                        errors, "Permissions", row_number, "path",
                        "HDFS path 不能与 database/table/column 同行填写",
                    )
                lower_path = path.lower()
                if lower_path.startswith("file:"):
                    _add_issue(
                        errors, "Permissions", row_number, "path",
                        "file: 是本地文件 URI，Guardian/TDFS 不管理该资源",
                    )
                elif path.upper() == "GLOBAL":
                    resource = ResourcePath(service_type=ServiceType.HDFS, path="GLOBAL")
                elif path.startswith("/"):
                    resource = ResourcePath(service_type=ServiceType.HDFS, path=path)
                elif lower_path.startswith("hdfs://"):
                    resource = ResourcePath(
                        service_type=ServiceType.HDFS,
                        path="hdfs://" + path[7:],
                    )
                else:
                    _add_issue(
                        errors, "Permissions", row_number, "path",
                        "HDFS 路径必须以 / 或 hdfs:// 开头，也可以填写 GLOBAL",
                    )
                allowed_actions = HDFS_ACTIONS
                scope = "HDFS_PATH"
            elif database:
                if column and not table:
                    _add_issue(
                        errors, "Permissions", row_number, "column",
                        "填写 column 时必须同时填写 table",
                    )
                normalized_database = "GLOBAL" if database.upper() == "GLOBAL" else database
                if table:
                    if normalized_database in {"*", "GLOBAL"}:
                        _add_issue(
                            errors, "Permissions", row_number, "database",
                            "全局 database 不能再填写 table 或 column",
                        )
                    if table == "*" or table.upper() == "GLOBAL":
                        _add_issue(
                            errors, "Permissions", row_number, "table",
                            "table 不允许 * 或 GLOBAL",
                        )
                    if column and (column == "*" or column.upper() == "GLOBAL"):
                        _add_issue(
                            errors, "Permissions", row_number, "column",
                            "column 不允许 * 或 GLOBAL",
                        )
                    resource = ResourcePath(
                        service_type=ServiceType.HIVE,
                        database=normalized_database,
                        table=table,
                        column=column or None,
                    )
                    allowed_actions = TABLE_COLUMN_ACTIONS
                    scope = "COLUMN" if column else "TABLE"
                else:
                    resource = ResourcePath(
                        service_type=ServiceType.HIVE,
                        database=normalized_database,
                    )
                    allowed_actions = DATABASE_ACTIONS
                    scope = "DATABASE"
            else:
                if table or column:
                    _add_issue(
                        errors, "Permissions", row_number, "database",
                        "填写 table/column 时必须填写 database",
                    )
                else:
                    _add_issue(
                        errors, "Permissions", row_number, "resource",
                        "database/table/column/path 不能全部为空",
                    )

            raw_actions = _parse_multi(
                action_text,
                sheet="Permissions",
                row=row_number,
                field="actions",
                errors=errors,
            )
            actions: list[str] = []
            for raw_action in raw_actions:
                action = raw_action.upper()
                if action == "OWNER":
                    _add_issue(errors, "Permissions", row_number, "actions", "不允许 OWNER")
                    continue
                if action == "ALL":
                    actions.extend(allowed_actions)
                elif action not in allowed_actions:
                    _add_issue(
                        errors, "Permissions", row_number, "actions",
                        f"{scope or '未知资源'} 不支持动作 {action}",
                    )
                else:
                    actions.append(action)
            actions = list(dict.fromkeys(actions))

            if resource is None or not actions:
                continue
            principal = Principal(principal_name, PrincipalType(principal_type))
            permissions: list[PermissionEntry] = []
            data_source = tuple(resource.to_guardian_data_source())
            for action in actions:
                key = (principal_type, principal_name, action, data_source)
                if key in seen_permissions:
                    continue
                seen_permissions.add(key)
                permissions.append(PermissionEntry(
                    action=action,
                    resource=resource,
                    principal=principal,
                    grantable=False,
                    heritable=True,
                    administrative=True,
                ))
            if permissions:
                plan.policies.append(Policy(
                    source="guardian-excel",
                    service_type=resource.service_type,
                    service_name="guardian-excel",
                    resources=[resource],
                    permissions=permissions,
                    description=f"Permissions row {row_number}: {principal_type} {principal_name}",
                ))

    workbook.close()
    permission_count = sum(len(policy.permissions) for policy in plan.policies)
    plan.source_metadata.update({
        "sheet_row_counts": row_counts,
        "permission_count": permission_count,
        "validation_warning_count": len(warnings),
    })
    report = {
        "status": "errors" if errors else "valid",
        "source": os.path.abspath(filepath),
        "error_count": len(errors),
        "warning_count": len(warnings),
        "errors": errors,
        "warnings": warnings,
        "summary": {
            "users": len(plan.users),
            "groups": len(plan.groups),
            "roles": len(plan.roles),
            "permissions": permission_count,
            "sheet_rows": row_counts,
        },
    }
    report_path = _write_report(validation_report, report) if validation_report else None
    if errors:
        raise ExcelValidationError(errors, report_path)
    if report_path:
        plan.source_metadata["validation_report"] = report_path
    return plan
